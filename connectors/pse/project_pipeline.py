"""Bulk documentary import; grouping source labels is not canonical station matching."""
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
import hashlib
import re
from openpyxl import load_workbook

STATUS = {'WNIOSEK w weryfikacji', 'WNIOSEK niekompletny',
          'WNIOSEK kompletny - w trakcie analizy technicznej i ekonomicznej',
          'WARUNKI PRZYŁĄCZENIA wydane', 'ODMOWA PRZYŁĄCZENIA',
          'UMOWA O PRZYŁĄCZENIE obowiązująca'}
HEADERS = {'B4':'Nazwa Obiektu', 'C4':'Nazwa OSD/OSP', 'D4':'Nazwa podmiotu',
           'E4':'Lokalizacja miejsca', 'F4':'Poziom napięcia', 'G4':'Wiele miejsc',
           'H4':'Rodzaj', 'I4':'Wprowadzana', 'J4':'Pobierana', 'Q4':'STATUS'}


def power(value) -> tuple[str | None, str | None]:
    if value is None or value == '-':
        return None, None
    text = str(value).strip().replace(',', '.')
    if isinstance(value, bool) or not re.fullmatch(r'\d+(?:\.\d+)?', text):
        return None, 'POWER_ANNOTATION_OR_INVALID'
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None, 'POWER_ANNOTATION_OR_INVALID'
    return str(number), None


def normalize_row(row: tuple, index: int) -> dict:
    reasons = []
    name, voltage, multi = row[4], row[5], row[6]
    if row[2] != 'PSE S.A.': reasons.append('OPERATOR_UNRECOGNIZED')
    if row[16] not in STATUS: reasons.append('STATUS_UNRECOGNIZED')
    if multi != 'NIE': reasons.append('MULTIPLE_POINTS_OR_UNKNOWN_FLAG')
    if type(voltage) not in (int, float) or voltage not in (110, 220, 400):
        reasons.append('VOLTAGE_REQUIRES_REVIEW')
    if not isinstance(name, str) or not name.strip() or any(c in name for c in '\n;'):
        reasons.append('POINT_LABEL_REQUIRES_REVIEW')
    export, e = power(row[8]); imp, i = power(row[9])
    if e: reasons.append('EXPORT_'+e)
    if i: reasons.append('IMPORT_'+i)
    label = name.strip() if isinstance(name, str) else None
    # No case folding, alias expansion, fuzzy merging or semicolon splitting.
    group_id = None
    association_issues = [x for x in reasons if not x.startswith(('EXPORT_', 'IMPORT_'))]
    if not association_issues:
        group_id = hashlib.sha256(f'PSE_PIPELINE|{label}|{int(voltage)}'.encode()).hexdigest()
    return {'record_id': f'pse_pipeline_row_{index}', 'row': index,
            'project_name': row[1], 'applicant': row[3], 'operator_reported': row[2],
            'point_reported': name, 'voltage_reported': voltage, 'multiple_points_reported': multi,
            'technology_reported': row[7], 'status_reported': row[16],
            'export_MW': export, 'import_MW': imp,
            'raw_export': row[8], 'raw_import': row[9],
            'profile_id': group_id, 'canonical_station_id': None,
            'classification': 'REPORTED', 'review_reasons': reasons}


def parse(path) -> dict:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book['Wykaz wspólny']
        date = re.search(r'Stan na (\d{2}\.\d{2}\.\d{4})', str(sheet['A1'].value))
        if not date: raise ValueError('SOURCE_DATE_MISSING')
        for cell, prefix in HEADERS.items():
            if not str(sheet[cell].value).startswith(prefix): raise ValueError('HEADER_CHANGED:'+cell)
        for cell in ('I4', 'J4'):
            if '[MW]' not in str(sheet[cell].value): raise ValueError('UNIT_CHANGED:'+cell)
        records, notes, blank = [], [], 0
        for index, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), 5):
            if all(v is None for v in row):
                blank += 1; continue
            if isinstance(row[0], str) and re.fullmatch(r'[A-ZŁ]\d?-', row[0]) and row[1] and all(v is None for v in row[2:]):
                notes.append({'row': index, 'marker': row[0], 'meaning': row[1]}); continue
            records.append(normalize_row(row, index))
        return {'source_date': datetime.strptime(date[1], '%d.%m.%Y').date().isoformat(),
                'records': records, 'footnotes': notes, 'blank_rows': blank,
                'physical_data_rows': sheet.max_row-4}
    finally:
        book.close()


def profiles(imported: dict, sample_size: int = 50) -> dict:
    groups = {}
    for r in imported['records']:
        if r['profile_id'] is None: continue
        g = groups.setdefault(r['profile_id'], {'profile_id':r['profile_id'],
            'name_reported': r['point_reported'].strip(), 'voltage_kV':r['voltage_reported'],
            'canonical_station_id': None, 'record_ids':[], 'status_counts':Counter(), 'review_record_ids':[]})
        g['record_ids'].append(r['record_id']);g['status_counts'][r['status_reported']]+=1
        if r['review_reasons']:g['review_record_ids'].append(r['record_id'])
    # Round-robin by voltage then source label; reproducible, not statistically representative.
    buckets = {v: sorted([g for g in groups.values() if g['voltage_kV']==v], key=lambda g:g['name_reported']) for v in (110,220,400)}
    selected=[]
    while len(selected)<min(sample_size,len(groups)):
        for bucket in buckets.values():
            if bucket and len(selected)<sample_size:selected.append(bucket.pop(0))
    all_rows=imported['records']
    return {'method':'pse_bulk_profiles_v1', 'source_date':imported['source_date'],
            'selection':'Round-robin 110/220/400 kV; alphabetical source labels within each bucket.',
            'metrics':{'source_records':len(all_rows), 'source_label_voltage_groups':len(groups),
                'grouped_records':sum(r['profile_id'] is not None for r in all_rows),
                'ungrouped_records':sum(r['profile_id'] is None for r in all_rows),
                'records_requiring_any_review':sum(bool(r['review_reasons']) for r in all_rows),
                'reason_counts':dict(Counter(x for r in all_rows for x in r['review_reasons'])),
                'selected_profiles':len(selected), 'validated_canonical_stations':0,
                'matching_error_rate':None, 'manual_review_minutes':None},
            'profiles':selected,
            'review_queue':[{'record_id':r['record_id'], 'row':r['row'], 'point_reported':r['point_reported'], 'reasons':r['review_reasons']} for r in all_rows if r['review_reasons']],
            'limitations':['Source label/voltage groups, not verified unique physical stations.',
                'Single PSE publication; no complete DSO or nationwide station coverage.',
                'Row counts are not unique projects, MW reservations, loading or adverse factors.',
                'No inference of operational status from agreements; annotated powers remain null.',
                'No manual ground-truth audit yet; error rate and manual effort unknown.']}
