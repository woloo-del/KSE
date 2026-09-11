"""Extract an auditable research sample, not a complete operational pipeline."""
import hashlib
import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def extract() -> dict:
    catalog = json.loads((ROOT/'data/catalog/data_sources.json').read_text(encoding='utf-8'))
    source = next(s for s in catalog['sources'] if s['source_id'] == 'PSE_PIPELINE')
    evidence = next(e for e in source['evidence'] if str(e['local_path']).endswith('.xlsx'))
    path = ROOT/evidence['local_path']
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != evidence['sha256']:
        raise ValueError('SNAPSHOT_HASH_MISMATCH')
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook['Wykaz wspólny']
    if 'Stan na 31.07.2026' not in str(sheet['A1'].value):
        raise ValueError('SOURCE_DATE_CHANGED')
    expected = {'B4':'Nazwa Obiektu', 'I4':'Wprowadzana', 'J4':'Pobierana',
                'Q4':'STATUS', 'X4':'Data zawarcia umowy'}
    for cell, prefix in expected.items():
        if not str(sheet[cell].value).startswith(prefix):
            raise ValueError('HEADER_CHANGED:'+cell)
    records = []
    for index, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), 5):
        if str(row[4]).strip().casefold() != 'radkowice':
            continue
        records.append({
            'research_record_id': f'pse_20260731_row_{index}',
            'project_name': row[1], 'applicant': row[3], 'connection_point_reported': row[4],
            'voltage_kV': row[5], 'technology_reported': row[7],
            'connection_export_MW': row[8], 'connection_import_MW': row[9],
            'status_reported': row[16],
            'connection_agreement_date': row[23].date().isoformat() if isinstance(row[23], datetime) else None,
            'delivery_start_date_reported': row[24].date().isoformat() if isinstance(row[24], datetime) else None,
            'classification': 'REPORTED', 'source_id': 'PSE_PIPELINE',
            'source_url': evidence['url'], 'source_date': '2026-07-31',
            'retrieval_date': evidence['retrieval_date'], 'snapshot_sha256': digest,
            'sheet': sheet.title, 'row': index,
            'cells': {'project_name':f'B{index}', 'applicant':f'D{index}',
                      'connection_point_reported':f'E{index}', 'voltage_kV':f'F{index}',
                      'connection_export_MW':f'I{index}', 'connection_import_MW':f'J{index}',
                      'status_reported':f'Q{index}', 'connection_agreement_date':f'X{index}',
                      'delivery_start_date_reported':f'Y{index}'},
        })
    workbook.close()
    return {'method':'radkowice_research_extract_v1', 'coverage':'Exact station-name matches in one PSE snapshot; not all PGE/PSE projects.',
            'warning':'Agreement and future delivery date do not prove physical connection. No node capacity or loading calculation.',
            'records':records}


if __name__ == '__main__':
    out = ROOT/'data/reference/radkowice_pse_projects_2026-07-31.json'
    out.write_text(json.dumps(extract(), ensure_ascii=False, indent=2)+'\n', encoding='utf-8', newline='\n')
    print(out)
