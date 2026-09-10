"""Offline integrity and semantic spot checks for the dated research snapshot.

This is not a production operator parser. It intentionally records source warnings
separately from validation failures and never rewrites downloaded evidence.
"""

from __future__ import annotations

import ast
import hashlib
import json
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import date, datetime, timezone
from importlib.metadata import version
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlparse

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw/research/2026-09-10"
REQUIRED = {
    "source_id", "operator", "source_name", "url", "data_owner", "country",
    "data_category", "grid_voltage", "geographical_scope", "data_format",
    "api_available", "gis_available", "update_frequency", "historical_data",
    "authentication", "license", "commercial_use", "reliability",
    "machine_readable", "scraping_required", "key_fields", "potential_application",
    "known_limitations", "data_quality_score", "priority", "last_verified",
    "last_successful_verification", "verification_status", "retrieval_date",
}
STATUSES = {"SAMPLE_VERIFIED", "CONTENT_REVIEWED", "DOCUMENTATION_REVIEWED", "DISCOVERED", "BLOCKED"}


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate() -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[dict[str, Any]] = []
    checks: list[dict[str, Any]] = []

    def check(name: str, condition: bool, detail: Any = None) -> None:
        checks.append({"name": name, "passed": bool(condition), "detail": detail})
        if not condition:
            errors.append(name)

    catalog = load_json(ROOT / "data/catalog/data_sources.json")
    sources = catalog["sources"]
    source_ids = [row["source_id"] for row in sources]
    check("catalog_unique_ids", len(source_ids) == len(set(source_ids)))
    for row in sources:
        sid = row["source_id"]
        check(f"catalog_fields:{sid}", REQUIRED <= row.keys(), sorted(REQUIRED - row.keys()))
        check(f"catalog_url:{sid}", urlparse(row["url"]).scheme in {"https", "http"} and bool(urlparse(row["url"]).netloc))
        check(f"catalog_status:{sid}", row["verification_status"] in STATUSES)
        check(f"no_uncalibrated_numeric_quality:{sid}", row["data_quality_score"] is None)
        check(f"source_rank:{sid}", row["reliability"] in "ABCDEFGH")
        for key in ["last_verified", "last_successful_verification", "source_date", "publication_date"]:
            value = row.get(key)
            if value:
                try:
                    date.fromisoformat(value)
                except ValueError:
                    check(f"date_format:{sid}:{key}", False, value)
        if row["verification_status"] in {"BLOCKED", "DISCOVERED"}:
            check(f"unverified_not_successful:{sid}", row["last_successful_verification"] is None)
        if row["verification_status"] == "SAMPLE_VERIFIED":
            check(f"sample_has_local_evidence:{sid}", any(p["local_path"] for p in row["evidence"]))

    probes = []
    for path in sorted((ROOT / "data/catalog").glob("probe_results_*.json")):
        probes.extend(load_json(path))
    probe_ids = {probe["source_id"] for probe in probes}
    check("probe_ids_unique", len(probe_ids) == len(probes))
    for row in sources:
        check(f"probe_references:{row['source_id']}", set(row["probe_ids"]) <= probe_ids)
    local_paths: set[Path] = set()
    for probe in probes:
        sid = probe["source_id"]
        if not probe["local_path"]:
            warnings.append({"source_id": sid, "code": "FETCH_ERROR", "http_status": probe["http_status"], "message": "Nie pobrano pliku; zachowano nieudaną próbę."})
            continue
        path = (ROOT / probe["local_path"]).resolve()
        check(f"snapshot_within_raw:{sid}", path.is_relative_to(RAW))
        check(f"snapshot_exists:{sid}", path.is_file())
        if path.is_file():
            local_paths.add(path)
            check(f"snapshot_bytes:{sid}", path.stat().st_size == probe["bytes"])
            check(f"snapshot_sha256:{sid}", file_hash(path) == probe["sha256"])
    check("all_raw_files_have_manifests", set(RAW.iterdir()) == local_paths)

    format_counts: Counter[str] = Counter()
    pdf_details = []
    pdf_first_pages: dict[str, str] = {}
    for path in sorted(local_paths):
        suffix = path.suffix.lower()
        format_counts[suffix] += 1
        try:
            if suffix == ".json":
                load_json(path)
            elif suffix == ".xml":
                root = ET.parse(path).getroot()
                check(f"wfs_capabilities:{path.name}", root.tag.endswith("WFS_Capabilities"), root.attrib.get("version"))
            elif suffix == ".pdf":
                reader = PdfReader(path)
                text = reader.pages[0].extract_text() or ""
                pdf_first_pages[path.name] = text
                pdf_details.append({"file": path.relative_to(ROOT).as_posix(), "pages": len(reader.pages), "first_page_text_chars": len(text)})
                check(f"pdf_text_layer:{path.name}", len(text.strip()) > 20)
            elif suffix == ".xlsx":
                workbook = load_workbook(path, read_only=True, data_only=False)
                check("pse_xlsx_expected_sheet", "Wykaz wspólny" in workbook.sheetnames)
                sheet = workbook["Wykaz wspólny"]
                head = " ".join(str(v) for row in sheet.iter_rows(min_row=1, max_row=4, values_only=True) for v in row if v is not None)
                check("pse_xlsx_snapshot_date", "31.07.2026" in head, "Header date, not retrieval date")
                check("pse_xlsx_direction_fields", "wprowadzana" in head.lower() and "pobierana" in head.lower())
                workbook.close()
            check(f"format_readable:{path.name}", True)
        except Exception as exc:
            check(f"format_readable:{path.name}", False, type(exc).__name__)

    pge = (RAW / "pge_capacity_landing.html").read_text(encoding="utf-8-sig")
    check("pge_block_body_identified", "odrzucone" in pge.lower())
    check("pge_catalog_not_data_success", next(r for r in sources if r["source_id"] == "PGE_LANDING")["verification_status"] == "BLOCKED")
    warnings.append({"source_id": "PGE_LANDING", "code": "SOURCE_ACCESS_BLOCKED", "http_status": 200, "message": "Odpowiedź HTTP 200 jest stroną odrzucenia, nie danymi operatora."})

    energa_header = pdf_first_pages.get("energa_pipeline_2026-08-31.pdf", "")
    check("energa_header_june_date", bool(re.search(r"30\s*\.\s*06\s*\.\s*2026", energa_header)))
    energa = next(r for r in sources if r["source_id"] == "ENERGA_PIPELINE")
    check("energa_conflict_preserved", energa["source_date"] is None and energa["date_conflict"]["resolution"] == "UNRESOLVED")
    warnings.append({"source_id": "ENERGA_PIPELINE", "code": "TEMPORAL_CONFLICT", "message": "Nazwa pliku 31.08.2026; pierwsza strona 30.06.2026. Nie rozstrzygnięto daty stanu."})

    nodes = load_json(RAW / "pse_nodes_sample.json")
    check("pse_nodes_shape", len(nodes["value"]) == 2 and all({"station_name", "switching_station_voltage", "busbar_code", "business_date"} <= row.keys() for row in nodes["value"]))
    check("pse_nodes_pagination", bool(nodes.get("nextLink")))
    redispatch = load_json(RAW / "pse_redispatch_sample.json")
    required = {"pv_red_balance", "wi_red_balance", "pv_red_network", "wi_red_network"}
    check("pse_redispatch_causes", all(required <= row.keys() for row in redispatch["value"]))
    check("pse_redispatch_null_preserved", all(row[key] is None for row in redispatch["value"] for key in required))
    flows = load_json(RAW / "pse_flows_sample.json")
    check("pse_flows_sections", all("section_code" in row for row in flows["value"]))
    charts = load_json(RAW / "energy_charts_pl_sample.json")
    series_ids = {item["id"] for item in charts.get("series", [])}
    observations = charts.get("data", [])
    check("energy_charts_v2_timeseries_shape", charts.get("schema_version") == "2.0" and bool(series_ids) and bool(observations))
    check("energy_charts_v2_series_keys", all(set(item.get("values", {})) == series_ids and "timestamp" in item for item in observations))
    check("energy_charts_v2_country_unit", charts.get("country") == "pl" and charts.get("unit") == "MW")
    check("energy_charts_v2_license_metadata", "CC BY 4.0" in charts.get("license", ""))
    warnings.append({"source_id": "PSE_API", "code": "SAMPLE_NOT_CURRENT_SNAPSHOT", "message": "Zapytania first=2 służą kontroli interfejsu; zwróciły starsze daty. Braki redysponowania nie są zerami."})

    for name in ["pse_openapi.json", "energy_charts_openapi.json"]:
        schema = load_json(RAW / name)
        check(f"openapi_contract:{name}", "paths" in schema and "info" in schema and "openapi" in schema)

    doc_paths = [ROOT / "README.md", *sorted((ROOT / "docs").glob("*.md"))]
    local_link_count = 0
    for path in doc_paths:
        text = path.read_text(encoding="utf-8")
        for target in re.findall(r"\]\(([^)]+)\)", text):
            target = target.strip("<>")
            if urlparse(target).scheme or target.startswith("#"):
                continue
            destination = (path.parent / unquote(target.split("#", 1)[0])).resolve()
            local_link_count += 1
            if destination == ROOT / "data/catalog/research_validation.json":
                continue  # This invocation creates the report after validation.
            check(f"local_link:{path.name}:{target}", destination.exists())
    for path in (ROOT / "scripts").glob("*.py"):
        ast.parse(path.read_text(encoding="utf-8"), filename=path.name)
        check(f"python_syntax:{path.name}", True)

    return {
        "checked_at_utc": datetime.now(timezone.utc).isoformat(),
        "scope": "Offline research catalog, raw integrity, format checks and selected semantic checks; not full operator parser validation",
        "status": "PASS_WITH_SOURCE_WARNINGS" if not errors else "FAIL",
        "source_count": len(sources),
        "source_status_counts": dict(Counter(row["verification_status"] for row in sources)),
        "probe_count": len(probes),
        "http_status_counts": dict(Counter(str(probe["http_status"]) for probe in probes)),
        "raw_file_count": len(local_paths),
        "format_counts": dict(format_counts),
        "pdf_details": pdf_details,
        "local_links_checked": local_link_count,
        "dependencies": {"python": sys.version.split()[0], "pypdf": version("pypdf"), "openpyxl": version("openpyxl")},
        "checks_passed": sum(item["passed"] for item in checks),
        "checks_failed": len(errors),
        "errors": errors,
        "source_warnings": warnings,
        "checks": checks,
    }


if __name__ == "__main__":
    report = validate()
    (ROOT / "data/catalog/research_validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({key: value for key, value in report.items() if key not in {"checks", "pdf_details"}}, ensure_ascii=False, indent=2))
    sys.exit(1 if report["errors"] else 0)
