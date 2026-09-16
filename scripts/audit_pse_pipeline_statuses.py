"""Audit statuses in one preserved PSE workbook, without inferring project identity."""
from __future__ import annotations

from collections import defaultdict
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def audit() -> dict:
    catalog = json.loads((ROOT / 'data/catalog/data_sources.json').read_text(encoding='utf-8'))
    source = next(s for s in catalog['sources'] if s['source_id'] == 'PSE_PIPELINE')
    evidence = next(e for e in source['evidence'] if str(e['local_path']).endswith('.xlsx'))
    path = ROOT / evidence['local_path']
    if hashlib.sha256(path.read_bytes()).hexdigest() != evidence['sha256']:
        raise ValueError('SNAPSHOT_HASH_MISMATCH')
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook['Wykaz wspólny']
        if 'Stan na 31.07.2026' not in str(sheet['A1'].value):
            raise ValueError('SOURCE_DATE_CHANGED')
        for cell, prefix in {'B4': 'Nazwa Obiektu', 'E4': 'Lokalizacja miejsca',
                             'Q4': 'STATUS', 'R4': 'Data złożenia wniosku'}.items():
            if not str(sheet[cell].value).startswith(prefix):
                raise ValueError('HEADER_CHANGED:' + cell)
        statuses = defaultdict(list)
        unclassified_rows = []
        station_mentions = []
        for index, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), 5):
            if not any(value is not None for value in row):
                continue
            status = row[16]
            if status is None or not str(status).strip():
                unclassified_rows.append(index)
            else:
                statuses[str(status)].append(index)
            if any('radkow' in str(value).casefold() for value in row):
                station_mentions.append({'row': index, 'point_reported': row[4],
                                         'status_reported': status, 'operator_object_id': row[0]})
        return {
            'method': 'pse_status_audit_v1', 'source_date': '2026-07-31',
            'classification': 'CALCULATED', 'formula': 'COUNT(source rows grouped by exact Q cell text)',
            'source_id': 'PSE_PIPELINE', 'source_url': evidence['url'],
            'retrieval_date': evidence['retrieval_date'], 'snapshot_sha256': evidence['sha256'],
            'sheet': sheet.title,
            'statuses': {key: {'row_count': len(rows), 'rows': rows} for key, rows in sorted(statuses.items())},
            'nonempty_rows_without_status': unclassified_rows,
            'radkow_substring_mentions': station_mentions,
            'limitations': [
                'Counts are source rows, not deduplicated projects or a complete nationwide inventory.',
                'Verification, incomplete application and technical/economic analysis are distinct stages.',
                'No automatic assignment to unanswered-request group C or to a shared bridge.',
                'Substring search is discovery only, not entity resolution; no matches does not prove absence.',
                'Rows without status are excluded from status counts, not assumed to be applications.',
                'The operator ID placeholder is not a usable canonical project identifier.',
            ],
        }
    finally:
        workbook.close()


if __name__ == '__main__':
    result = audit()
    result['code_sha256'] = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    target = ROOT / 'data/reference/pse_status_audit_2026-07-31_v1.json'
    content = (json.dumps(result, ensure_ascii=False, indent=2) + '\n').encode('utf-8')
    if target.exists():
        if target.read_bytes() != content:
            raise ValueError('EXISTING_DIFFERENT_AUDIT')
    else:
        with target.open('xb') as stream:
            stream.write(content)
    print(target)
