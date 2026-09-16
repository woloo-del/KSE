"""Compare attribute multisets in supplied XLSX/GPKG exports; no entity resolution."""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from connectors.gis.geopackage import read_features

# WPZP repeats one display header for two different source fields. Map by position.
WPZP_FIELDS=['Name','OSD','Status','W-PZP','Link','Year','More accurate location','Comment','Voivodship','District','Commune']
WPZP_HEADERS=['Nazwa obiektu','Nazwa Operatora ze względu na położenie','Status','Województwo','Link do dokumentu źródłowego','Rok sporządzenia dokumentu źródłowego','Identyfikacja dokładniejszego położenia','Komentarz','Województwo','Powiat','Gmina']
ALIASES={'(1) Link':'(1) Link do źródła','(2) Link':'(2) Link do źródła'}


def normalized(value: object) -> str | None:
    # Preserve internal whitespace, punctuation and wording. No fuzzy matching.
    return str(value).replace('\r\n','\n').strip() if value is not None else None


def read_table(path: Path) -> tuple[list[str], list[dict]]:
    workbook=load_workbook(path,read_only=True,data_only=False)
    try:
        if len(workbook.worksheets)!=1: raise ValueError('EXPECTED_SINGLE_WORKSHEET')
        sheet=workbook.worksheets[0]
        rows=sheet.iter_rows(values_only=True)
        header=list(next(rows))
        if header[0] is not None: raise ValueError('UNEXPECTED_FIRST_COLUMN')
        headers=header[1:]
        if path.stem.endswith('WPZP'):
            if headers!=WPZP_HEADERS: raise ValueError('WPZP_SCHEMA_CHANGED')
            fields=WPZP_FIELDS
        else:
            if any(not isinstance(h,str) for h in headers): raise ValueError('INVALID_HEADER')
            fields=[ALIASES.get(h,h) for h in headers]
            if len(set(fields))!=len(fields): raise ValueError('DUPLICATE_HEADERS')
        result=[]
        for number,row in enumerate(rows,2):
            if not any(v is not None for v in row): continue
            if row[0] is not None: raise ValueError('UNEXPECTED_FIRST_COLUMN_VALUE')
            if any(isinstance(v,str) and v.startswith('=') for v in row): raise ValueError('FORMULA_REQUIRES_REVIEW')
            result.append({'row':number,'values':[normalized(v) for v in row[1:]]})
        return fields,result
    finally:
        workbook.close()


def compare_rows(fields: list[str], rows: list[dict], features: list) -> dict:
    available=defaultdict(list)
    for feature in features:
        if not all(field in feature.properties for field in fields): raise ValueError('GPKG_SCHEMA_MISMATCH')
        key=tuple(normalized(feature.properties[field]) for field in fields)
        available[key].append(feature.fid)
    matches,unmatched=[],[]
    for row in rows:
        key=tuple(row['values'])
        if available.get(key):
            matches.append({'xlsx_row':row['row'],'gpkg_fid':available[key].pop(0)})
        else: unmatched.append(row)
    extras=[{'gpkg_fid':fid,'values':list(key)} for key,ids in available.items() for fid in ids]
    return {'matches':matches,'unmatched_xlsx_rows':unmatched,'unmatched_gpkg_rows':extras}


def compare(root: Path, output: Path) -> dict:
    private=(ROOT/'data/private').resolve()
    root,output=root.resolve(),output.resolve()
    if not root.is_relative_to(private) or not output.is_relative_to(private) or output.is_relative_to(root) or output.exists():
        raise ValueError('Use private source directory and a new separate private output directory.')
    paths=sorted(root.glob('*.xlsx'))
    if not paths: raise ValueError('NO_XLSX_FILES_FOUND')
    results,inputs=[],[]
    for path in paths:
        pair=path.with_name(path.stem.removesuffix(' - tabela')+'.gpkg')
        hashes={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in [path,pair]}
        inputs.extend({'path':p.relative_to(private).as_posix(),'sha256':digest} for p,digest in hashes.items())
        fields,rows=read_table(path)
        features,errors=read_features(pair)
        result=compare_rows(fields,rows,features)
        pair_fields=('Nazwa zadania inwestycyjnego','Podstawowy cel realizacji zadania inwestycyjnego')
        if all(field in fields for field in pair_fields) and result['unmatched_xlsx_rows']:
            left,right=(fields.index(field) for field in pair_fields)
            swapped=[]
            for row in rows:
                values=list(row['values'])
                values[left],values[right]=values[right],values[left]
                swapped.append({'row':row['row'],'values':values})
            diagnostic=compare_rows(fields,swapped,features)
            result['column_swap_diagnostic']={'fields':pair_fields,'matches_after_swap':len(diagnostic['matches']),
                'source_modified':False,'status':'HYPOTHESIS_ONLY_ORIGINAL_MISMATCHES_RETAINED'}
        result.update({'xlsx':path.name,'gpkg':pair.name,'fields':fields,'xlsx_rows':len(rows),'gpkg_rows':len(features),'geometry_read_errors':errors})
        results.append(result)
        for p,digest in hashes.items():
            if hashlib.sha256(p.read_bytes()).hexdigest()!=digest: raise ValueError('SOURCE_CHANGED_DURING_COMPARISON')
    result={'method':'gis_table_comparison_v1','access':'PRIVATE','retrieval_date':datetime.now(timezone.utc).isoformat(),
        'inputs':inputs,'tables':results,'limitations':['Attribute multiset equivalence is not independent source corroboration.',
            'Row matches do not merge entities. Repeated rows are consumed one at a time.',
            'Does not establish current investment status, licensing or correctness of source wording.']}
    output.mkdir(parents=True)
    (output/'table_comparison.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf8')
    print(f"Compared {len(results)} tables; {sum(len(r['matches']) for r in results)} matching rows; {sum(len(r['unmatched_xlsx_rows']) for r in results)} unmatched XLSX rows.")
    return result


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('root',type=Path);parser.add_argument('output',type=Path)
    args=parser.parse_args();compare(args.root,args.output)
